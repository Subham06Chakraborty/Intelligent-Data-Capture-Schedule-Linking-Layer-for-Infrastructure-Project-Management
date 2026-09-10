import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    primary_color = RGBColor(0, 102, 204)    # #0066CC (Blue for Status Report)
    dark_heading = RGBColor(15, 23, 42)
    sub_heading = RGBColor(51, 65, 85)
    body_color = RGBColor(30, 41, 59)
    success_color = RGBColor(0, 153, 76)     # Green for Completed items

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("INTELLITRACK — PROJECT PROGRESS REPORT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Status of Implementation & Achieved Milestones")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = sub_heading

    def add_section_header(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = primary_color

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = dark_heading
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = body_color

    def add_bullet(bold_part, text_part, completed=False):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if completed:
            rc = p.add_run("[COMPLETED] ")
            rc.font.name = "Calibri"
            rc.font.size = Pt(10)
            rc.font.bold = True
            rc.font.color.rgb = success_color

        rb = p.add_run(bold_part)
        rb.font.name = "Calibri"
        rb.font.size = Pt(10.5)
        rb.font.bold = True
        rb.font.color.rgb = dark_heading
        rt = p.add_run(text_part)
        rt.font.name = "Calibri"
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = body_color

    # Add content
    add_section_header("1. Overall Status")
    add_body("Current Status: 100% Prototype Completion")
    add_body("The IntelliTrack prototype successfully bridges the gap between manual field reporting and Primavera L5/L6 schedules. It meets all technical requirements of the SIH Smart Automation problem statement, seamlessly connecting unstructured data ingestion, LLM validation, and ML delay forecasting.")

    add_section_header("2. Frontend Milestones (React + Vite)")
    add_bullet("Upload & Processing Dashboard: ", "Built an intuitive drag-and-drop interface mapping files (PDF, Excel, TXT) to the AI extraction pipeline.", completed=True)
    add_bullet("Supervisor Chat Interface: ", "Developed a WhatsApp-style interface with live voice recording (browser MediaRecorder) and audio file attachments for seamless field data capture.", completed=True)
    add_bullet("L5/L6 Activities Review Board: ", "Created a planner-focused dashboard for human-in-the-loop validation, allowing engineers to confirm or reject AI fuzzy matches.", completed=True)
    add_bullet("ML Analytics Dashboard: ", "Designed an interactive page with real-time model health metrics, Recharts-based delay forecasts, and feature importance breakdowns.", completed=True)
    add_bullet("Design System: ", "Applied a premium, Oil India-inspired dark glassmorphism aesthetic across all pages.", completed=True)

    add_section_header("3. Backend & AI Milestones (FastAPI + Groq + Scikit-Learn)")
    add_bullet("LLM Extraction Engine: ", "Integrated Groq's LLaMA-3 70B model with domain-specific prompt engineering to normalize colloquial construction terms into standard engineering disciplines.", completed=True)
    add_bullet("Voice Transcription Pipeline: ", "Wired Whisper-large-v3 into the /api/v1/ingest/voice endpoint to support real-time Hindi/English voice note translation and extraction.", completed=True)
    add_bullet("Semantic Schedule Matching: ", "Deployed SentenceTransformers ('all-MiniLM-L6-v2') to perform fuzzy semantic matching against the baseline schedule with high-accuracy confidence scores.", completed=True)
    add_bullet("ML Delay Predictors: ", "Successfully loaded Databricks-trained Random Forest and Gradient Boosting models into memory, writing a custom shim to resolve pickling incompatibilities on modern Python environments.", completed=True)
    add_bullet("Local Firestore Mock: ", "Ensured zero downtime by writing a local JSON-based fallback for Firestore, ensuring all extracted data and audit trails persist safely locally.", completed=True)

    add_section_header("4. Deliverables & Documentation")
    add_bullet("ML & Judge QA Guide: ", "Generated 'IntelliTrack_ML_and_Judge_QA_Guide.docx' containing technical defense points for the hackathon judges.", completed=True)
    add_bullet("Architecture & Flows: ", "Generated 'IntelliTrack_Architecture_and_Flows.docx' outlining the 5 main system workflows.", completed=True)
    add_bullet("Progress Report: ", "This document, summarizing all accomplished development sprints.", completed=True)

    add_section_header("5. Next Steps")
    add_body("The prototype is fully functional. Teams should now focus on:")
    add_bullet("", "Rehearsing the live demo (Uploading a sample report -> Chatting with the Supervisor Agent -> Reviewing the Match -> Seeing the ML Delay Risk update).")
    add_bullet("", "Reviewing the Judge QA document to comfortably defend the use of Databricks ML models and Groq LLMs.")

    output_path = r"d:\DATABRICKS\IntelliTrack_Progress_Report.docx"
    doc.save(output_path)
    print(f"SUCCESS: Document saved to {output_path}")

if __name__ == "__main__":
    create_document()
