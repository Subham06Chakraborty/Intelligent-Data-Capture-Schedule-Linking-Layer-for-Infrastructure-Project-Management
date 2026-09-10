import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    primary_color = RGBColor(255, 75, 0)
    dark_heading = RGBColor(15, 23, 42)
    sub_heading = RGBColor(51, 65, 85)
    body_color = RGBColor(30, 41, 59)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("INTELLITRACK — ARCHITECTURE & WORKFLOWS")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Detailed Documentation of System Flows & Implementations")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = sub_heading

    def add_section_header(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = primary_color

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = dark_heading
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = body_color

    def add_bullet(bold_part, text_part):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        rb = p.add_run(bold_part)
        rb.font.name = "Calibri"
        rb.font.size = Pt(10.5)
        rb.font.bold = True
        rb.font.color.rgb = dark_heading
        rt = p.add_run(text_part)
        rt.font.name = "Calibri"
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = body_color

    # Add content
    add_section_header("1. Executive Summary of Implementations")
    add_body("IntelliTrack was built to bridge the gap between planning and execution. We have covered the complete end-to-end flow of ingesting chaotic field data, parsing it with AI, matching it to strict L5/L6 schedules, validating it, and using it for ML delay predictions.")

    add_section_header("2. Workflow 1: Heterogeneous File Ingestion & Extraction")
    add_bullet("How we covered it: ", "Built an UploadPage in the frontend connected to an ingestion pipeline in the backend.")
    add_bullet("File Processing: ", "PDF, Excel, and Text files are parsed dynamically.")
    add_bullet("LLM Extraction: ", "Text is fed to Groq's LLaMA-3 70B model with a specialized prompt to extract activities, disciplines, start dates, and end dates as structured JSON.")

    add_section_header("3. Workflow 2: Supervisor Chat & Voice Interface ('Time Agent')")
    add_bullet("How we covered it: ", "Built a dedicated SupervisorChatPage to capture data directly from site engineers without complex forms.")
    add_bullet("Live Voice Recording: ", "Implemented browser-native MediaRecorder to capture audio directly on the page, mimicking WhatsApp.")
    add_bullet("Whisper Transcription: ", "Voice notes are transcribed in milliseconds using Groq's whisper-large-v3 model.")
    add_bullet("Conversational Agent: ", "LLaMA-3 manages context history to extract the 4 required fields from English or Hindi sentences seamlessly.")

    add_section_header("4. Workflow 3: Semantic Schedule Matching")
    add_bullet("How we covered it: ", "Implemented fuzzy_service.py to bridge the language gap between the field (e.g., 'pipe joint done') and Primavera (e.g., 'Erect Line 24').")
    add_bullet("Algorithm: ", "Using SentenceTransformers ('all-MiniLM-L6-v2') to convert text into embeddings, computing cosine similarity combined with RapidFuzz lexical matching.")
    add_bullet("Confidence Thresholds: ", "Matches are classified automatically. High confidence items are 'Matched', low confidence items are 'Flagged' for human review.")

    add_section_header("5. Workflow 4: Planner Review & Audit Trail")
    add_bullet("How we covered it: ", "Built an interactive L5/L6 Activities Log dashboard (ActivitiesPage.jsx).")
    add_bullet("Human-in-the-Loop Validation: ", "Planners can filter activities by 'Requires Review' and manually Confirm or Reject semantic matches.")
    add_bullet("Audit Trail: ", "Every extraction, AI match, and planner review is securely logged in the backend with timestamps, user IDs, and confidence scores.")

    add_section_header("6. Workflow 5: Real-Time Machine Learning & Analytics")
    add_bullet("How we covered it: ", "Integrated trained ML models built via Databricks Delta tables into a live FastAPI backend.")
    add_bullet("Delay Prediction: ", "As soon as an activity is logged, a Random Forest Classifier and Gradient Boosting Regressor predict the delay risk and forecasted duration based on disciplines, contractor scores, and monsoon flags.")
    add_bullet("Institutional Memory: ", "The entire pipeline builds a structured, queryable actual-progress dataset in Firestore, preserving project knowledge for future planning.")

    output_path = r"d:\DATABRICKS\IntelliTrack_Architecture_and_Flows.docx"
    doc.save(output_path)
    print(f"SUCCESS: Document saved to {output_path}")

if __name__ == "__main__":
    create_document()
